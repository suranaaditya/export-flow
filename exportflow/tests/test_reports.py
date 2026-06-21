import frappe

try:
	from frappe.tests import IntegrationTestCase
except ImportError:  # frappe < 16
	from frappe.tests.utils import FrappeTestCase as IntegrationTestCase

from exportflow.reports import REPORTS, report_data, report_export, report_list
from exportflow.tests.test_dropship import _suffix, make_customer
from exportflow.tests.test_logistics import book_deal, make_plain_item

REPORT_KEYS = ["realization", "incentive", "sales_register", "gst_export", "merchanting", "mis"]


class TestReports(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.company = frappe.db.get_single_value("Global Defaults", "default_company")

	def test_catalog(self):
		keys = {r["key"] for r in report_list()}
		self.assertTrue(set(REPORT_KEYS) <= keys, f"missing: {set(REPORT_KEYS) - keys}")

	def test_data_shape(self):
		for key in REPORT_KEYS:
			d = report_data(key)
			self.assertTrue(d["title"], key)
			self.assertGreater(len(d["columns"]), 0, key)
			self.assertIsInstance(d["rows"], list)
			self.assertTrue(all("control" in f and "field" in f for f in d["filters"]), key)
			if d["rows"]:
				# the rows are dicts whose keys overlap the declared columns
				col_keys = {c["key"] for c in d["columns"]}
				self.assertTrue(col_keys & set(d["rows"][0]), key)

	def test_sales_register_rows(self):
		sfx = _suffix()
		item = make_plain_item(f"_Test EF Rep Item {sfx}")
		cust = make_customer(f"_Test EF Rep Cust {sfx}")
		book_deal(self.company, cust, [{"item_code": item, "qty": 5, "rate": 100}])
		rows = report_data("sales_register")["rows"]
		self.assertTrue(any(sfx in str(r.get("customer", "")) for r in rows), "the new SO is in the register")
		# INR value is the SO base grand total (positive)
		mine = next(r for r in rows if sfx in str(r.get("customer", "")))
		self.assertGreater(mine["inr_value"], 0)

	def test_export_xlsx(self):
		report_export("realization", "xlsx", rows="[]")
		self.assertEqual(frappe.local.response.type, "binary")
		self.assertTrue(frappe.local.response.filename.endswith(".xlsx"))
		self.assertTrue(frappe.local.response.filecontent, "xlsx bytes produced")

	def test_export_xlsx_professional_formatting(self):
		"""Feedback #25: the xlsx is professionally formatted — a bold+frozen header,
		auto-filter, real numbers with Indian grouping, and a bold totals row."""
		import io
		import json

		import openpyxl

		cols = REPORTS["realization"][0]["columns"]
		inr_idx = next((i for i, c in enumerate(cols, start=1) if c["type"] == "inr"), None)
		self.assertIsNotNone(inr_idx, "realization has a money column to total")

		def val(c):
			return 12345.67 if c["type"] == "inr" else (3 if c["type"] == "num" else "x")

		rows = [{c["key"]: val(c) for c in cols} for _ in range(2)]
		report_export("realization", "xlsx", rows=json.dumps(rows))
		ws = openpyxl.load_workbook(io.BytesIO(frappe.local.response.filecontent)).active

		self.assertTrue(ws.cell(row=1, column=1).font.bold, "styled header")
		self.assertEqual(ws.freeze_panes, "A2", "header row frozen")
		self.assertTrue(ws.auto_filter.ref, "auto-filter set")

		money = ws.cell(row=2, column=inr_idx)
		self.assertIsInstance(money.value, (int, float), "money is a real number, not text")
		self.assertIn(",", money.number_format)

		total_row = 1 + len(rows) + 1  # header + data + totals
		total_cell = ws.cell(row=total_row, column=inr_idx)
		self.assertTrue(total_cell.font.bold, "bold totals row")
		self.assertEqual(round(total_cell.value or 0, 2), 24691.34, "money column summed")

	def test_export_rows_are_formatted(self):
		# a single row round-trips into the xlsx without raising
		report_export(
			"sales_register",
			"xlsx",
			rows='[{"name":"SO-1","transaction_date":"2026-01-01","customer":"X","country":"India","currency":"USD","grand_total":100,"inr_value":8300,"incoterm":"FOB","status":"To Bill"}]',
		)
		self.assertTrue(frappe.local.response.filecontent)

	def test_export_columns_subset(self):
		# only the chosen columns (in canonical order) reach the file
		import io

		import openpyxl

		report_export("sales_register", "xlsx", rows="[]", columns='["customer", "name"]')
		ws = openpyxl.load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
		self.assertEqual([c.value for c in ws[1]], ["Sales order", "Customer"])

	def test_export_unknown_columns_fallback(self):
		import io

		import openpyxl

		# all-unknown selection falls back to the full column set
		report_export("sales_register", "xlsx", rows="[]", columns='["bogus", "nope"]')
		ws = openpyxl.load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
		self.assertEqual(len(list(ws[1])), len(REPORTS["sales_register"][0]["columns"]))
		# a mix keeps only the valid key
		report_export("sales_register", "xlsx", rows="[]", columns='["name", "bogus"]')
		ws2 = openpyxl.load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
		self.assertEqual([c.value for c in ws2[1]], ["Sales order"])

	def test_report_pdf_layout(self):
		"""_report_pdf fits the chosen columns: few -> portrait, many -> landscape,
		deselected labels absent, the filter summary stamped, page numbers in footer."""
		from exportflow.reports import _report_pdf

		cols = REPORTS["realization"][0]["columns"]
		few = [c for c in cols if c["key"] in ("export_invoice", "customer", "inr_value")]
		html, options = _report_pdf("Realization", "Status: Overdue", few, [], {"inr_value": 100.0})
		self.assertEqual(options["orientation"], "Portrait")
		self.assertIn("INR value", html)  # a chosen column's label
		self.assertNotIn("FIRC", html)  # a deselected column's label is gone
		self.assertIn("Status: Overdue", html)  # filter summary stamped under the title
		self.assertIn("<colgroup>", html)  # proportional widths => single page width
		self.assertIn("table-layout:fixed", html)
		self.assertIn("[page]", options["footer-right"])
		# the full column set is wide => landscape
		_html2, options2 = _report_pdf("Realization", "", cols, [], {})
		self.assertEqual(options2["orientation"], "Landscape")

	def test_report_pdf_width_driven_landscape(self):
		"""<=6 columns but wide (weights sum >70) still flips to landscape — pins the
		width term independently of the column-count term."""
		from exportflow.reports import _report_pdf

		cols = REPORTS["realization"][0]["columns"]
		wide6 = [c for c in cols if c["key"] in ("customer", "firc_no", "ebrc_number", "ad_bank", "export_invoice", "inr_value")]
		self.assertEqual(len(wide6), 6)
		_h, opts = _report_pdf("R", "", wide6, [], {"inr_value": 1.0})
		self.assertEqual(opts["orientation"], "Landscape")

	def test_report_pdf_totals_label_position(self):
		"""The 'Total' caption lands on the first NON-money column even when the leading
		shown column is itself a money column (and never crashes when all are money)."""
		from exportflow.reports import _report_pdf

		cols = REPORTS["realization"][0]["columns"]
		# leading shown column is money (inr_value) -> label must still render (on status)
		lead_money = [c for c in cols if c["key"] in ("inr_value", "status")]
		html, _o = _report_pdf("R", "", lead_money, [], {"inr_value": 100.0})
		self.assertIn("tlbl'>Total", html)
		# every shown column is money -> no caption cell, but must not raise (the .tlbl
		# CSS rule is always present, so assert against the rendered cell markup)
		all_money = [c for c in cols if c["key"] in ("inr_value", "realized_inr", "outstanding_inr")]
		h2, _o2 = _report_pdf("R", "", all_money, [], {c["key"]: 1.0 for c in all_money})
		self.assertNotIn("tlbl'>Total", h2)

	def test_export_scalar_columns_no_crash(self):
		"""A non-list `columns` payload is ignored (falls back to all columns), not a 500."""
		import io

		import openpyxl

		report_export("sales_register", "xlsx", rows="[]", columns="5")
		ws = openpyxl.load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
		self.assertEqual(len(list(ws[1])), len(REPORTS["sales_register"][0]["columns"]))
		report_export("sales_register", "xlsx", rows="[]", columns="true")
		self.assertTrue(frappe.local.response.filecontent)

	def test_export_malformed_rows(self):
		"""Malformed JSON rows -> clean ValidationError; wrong-shaped rows -> filtered, no crash."""
		with self.assertRaises(frappe.ValidationError):
			report_export("sales_register", "xlsx", rows="{")
		report_export("sales_register", "xlsx", rows="[1, 2, 3]")
		self.assertTrue(frappe.local.response.filecontent)

	def test_report_pdf_dynamic_page_size(self):
		"""Wide column sets step the page up (A4 → A0) instead of cramming into A4."""
		from exportflow.reports import _report_pdf

		# the full 78-column MIS set needs the largest page
		_h, opts = _report_pdf("MIS", "", REPORTS["mis"][0]["columns"], [], {})
		self.assertEqual(opts["page-size"], "A0")
		self.assertEqual(opts["orientation"], "Landscape")
		# a normal ~12-column report stays on A4
		_h2, opts2 = _report_pdf("MTT", "", REPORTS["merchanting"][0]["columns"], [], {})
		self.assertEqual(opts2["page-size"], "A4")

	def test_mis_register(self):
		"""The MIS workbook is the 78-column line-level register with NO totals row
		(per-shipment values repeat across lines, so a column sum would double-count)."""
		meta = REPORTS["mis"][0]
		self.assertEqual(len(meta["columns"]), 78)
		self.assertFalse(meta.get("totals", True))
		d = report_data("mis")
		self.assertEqual(d.get("totals"), False)
		self.assertIsInstance(d["rows"], list)
		# exports cleanly even though it has inr columns (fob/RoDTEP/DBK) — no totals row
		report_export("mis", "xlsx", rows=frappe.as_json(d["rows"]))
		self.assertTrue(frappe.local.response.filecontent)

	def test_unknown_report(self):
		with self.assertRaises(frappe.ValidationError):
			report_data("bogus")
		with self.assertRaises(frappe.ValidationError):
			report_export("bogus", "xlsx", rows="[]")
